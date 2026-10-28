"""Le classeur de suivi CFE : une synthèse d'une page, les dossiers comme dans l'application, puis le détail des avis."""

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from .notice_extraction import PAYMENT_LABELS, debited_account, review_label
from .situations import URGENCY, dossier_reading, dossier_situation, earliest_due_date, mail_text, situation_label

# The application's palette: ink on light grey, cobalt for links and headings.
INK = '131A2C'
MUTED = '59627A'
COBALT = '2F45D8'
ROW_LINE = 'EEF0F5'
WHITE = 'FFFFFF'
FONT = 'Aptos'
# One colour per situation, as on the application's chips.
SITUATION_COLORS = {
    'PAY': ('FFF0DC', 'A85600'),
    'DEBIT': ('DFF1F6', '1A6A80'),
    'NONE': ('EBEEF4', '59627A'),
    'REVIEW': ('FCE5EB', 'B01F48'),
    'RETRY': ('FCE5EB', 'B01F48'),
    'ABSENT': ('F0F2F6', '59627A'),
    'WAIT': ('F0F2F6', '59627A'),
    'MANUAL': ('FDEFD2', '8A5300'),
    'CHECKED': ('E3F4EA', '23704A'),
}
MAIL_COLORS = {'Prêt': ('E8EBFD', COBALT), 'À confirmer': ('FFF0DC', 'A85600'), 'Sans adresse': ('EBEEF4', MUTED)}
KIND_LABELS = {'ACOMPTE': 'Acompte', 'SOLDE': 'Solde', 'SUPPLEMENTAIRE': 'Rôle supplémentaire'}
SUMMARY_LABELS = {
    'RETRY': 'À relancer ou sans habilitation',
    'REVIEW': 'Lecture à vérifier',
    'WAIT': 'En attente de collecte',
    'PAY': 'À régler par le client',
    'DEBIT': 'Prélevé',
    'NONE': 'Rien à payer',
    'ABSENT': 'Sans avis',
    'MANUAL': 'À vérifier manuellement',
    'CHECKED': 'Vérifiés manuellement',
}
DOSSIER_SHEET = 'Dossiers'


def _text(value):
    # Always write externally supplied names and notes as literal text.
    return '' if value is None else str(value)


def _date(value, with_time=False):
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    for pattern in (('%Y-%m-%dT%H:%M:%S' if with_time else '%d/%m/%Y'), '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(str(value)[:19] if with_time else str(value), pattern)
        except ValueError:
            pass
    return _text(value)


def _font(size=10, bold=False, color=INK, **extra):
    return Font(name=FONT, size=size, bold=bold, color=color, **extra)


def _fill(color):
    return PatternFill('solid', fgColor=color)


def _link(cell, target, label):
    cell.value = label
    cell.hyperlink = Path(target).as_uri()
    cell.font = _font(bold=True, color=COBALT, underline='single')


def _plural(count, singular, plural):
    return f'{count} {singular if count == 1 else plural}'


def _sheet_heading(sheet, title, context, note):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.row_dimensions[1].height = 10
    sheet['A2'] = title
    sheet['A2'].font = _font(17, True)
    sheet.row_dimensions[2].height = 30
    sheet['A3'] = context
    sheet['A3'].data_type = 's'
    sheet['A3'].font = _font(color=MUTED)
    sheet['A4'] = note
    sheet['A4'].font = _font(color=MUTED, italic=True)
    sheet.row_dimensions[5].height = 6


def _table(sheet, headers, rows, widths, name, *, money=(), dates=(), datetimes=(), texts=(), wrap=(), links=None):
    """A filtrable table under the heading: one header row in ink, thin lines between rows."""
    links = links or {}
    for index, (header, width) in enumerate(zip(headers, widths), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(6, index, header)
        cell.fill = _fill(INK)
        cell.font = _font(bold=True, color=WHITE)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[6].height = 30
    for offset, values in enumerate(rows):
        row = 7 + offset
        sheet.row_dimensions[row].height = 34 if any(len(str(value or '')) > 70 for value in values) else 24
        for index, value in enumerate(values, 1):
            header = headers[index - 1]
            cell = sheet.cell(row, index, value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.font = _font()
            cell.border = Border(bottom=Side(style='thin', color=ROW_LINE))
            cell.alignment = Alignment(
                vertical='center', wrap_text=header in wrap, horizontal='right' if header in money else None
            )
            if header in money:
                cell.number_format = '#,##0.00 "€"'
            elif header in dates:
                cell.number_format = 'dd/mm/yyyy'
            elif header in datetimes:
                cell.number_format = 'dd/mm/yyyy hh:mm'
            if header in texts:
                cell.number_format = '@'
                cell.quotePrefix = True
                cell.font = _font(color='3A4258')
        for header, (targets, label) in links.items():
            target = targets[offset]
            if target:
                _link(sheet.cell(row, headers.index(header) + 1), target, label(target) if callable(label) else label)
    last = get_column_letter(len(headers))
    sheet.freeze_panes = 'C7'
    sheet.print_title_rows = '1:6'
    sheet.page_setup.orientation = 'landscape'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    # The table brings its own filter: a second filter of the sheet on the same cells makes Excel refuse the file.
    if not rows:
        sheet.auto_filter.ref = f'A6:{last}6'
    if rows:
        table = Table(displayName=name, ref=f'A6:{last}{6 + len(rows)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
        sheet.add_table(table)


def _chip(cell, colors):
    background, foreground = colors
    cell.fill = _fill(background)
    cell.font = _font(bold=True, color=foreground)


def dossier_rows(records):
    """Each dossier and year with its situation, sorted by urgency, as in the application's Dossiers tab."""
    rows = []
    for record in records:
        dossier, notices = record['dossier'], record['notices']
        situation = dossier_situation(dossier, notices)
        due, _mode, _reading = dossier_reading(notices)
        rows.append(
            {
                **record,
                'situation': situation,
                'due': due,
                'label': situation_label(dossier, situation),
                'mail': mail_text(dossier, situation),
            }
        )
    rows.sort(key=lambda row: (URGENCY.index(row['situation']), row['dossier']['name'].casefold(), row['year']))
    return rows


def _short_date(stamp):
    return f'{stamp[8:10]}/{stamp[5:7]}/{stamp[:4]}' if stamp else ''


def _information(row):
    dossier, situation = row['dossier'], row['situation']
    if situation == 'CHECKED':
        return 'Vérifié manuellement le ' + _short_date(dossier['verified_at']) + '.'
    if situation == 'MANUAL':
        return f'{dossier.get("failures") or 0} échecs : à vérifier sur impots.gouv.fr. ' + _text(dossier['error'])
    if dossier['error']:
        return _text(dossier['error'])
    if dossier['status'] == 'SANS_HABILITATION':
        return 'Sans habilitation · capture dans le dossier client.'
    if situation == 'REVIEW':
        return 'Montant à vérifier sur l’avis.'
    if situation == 'ABSENT':
        return 'Aucun avis · preuve dans le dossier client.'
    if situation == 'WAIT':
        return 'Pas encore collecté.'
    return ''


def _summary(sheet, campaign_name, scope, rows, cabinet):
    """One printable page: where the dossiers stand, the mails, what is left to do."""
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.showRowColHeaders = False
    sheet.sheet_view.zoomScale = 90
    for letter, width in {
        'A': 3,
        'B': 34,
        'C': 10,
        'D': 12,
        'E': 30,
        'F': 12,
        'G': 12,
        'H': 12,
        'I': 12,
        'J': 14,
    }.items():
        sheet.column_dimensions[letter].width = width
    sheet['B2'] = 'Suivi CFE'
    sheet['B2'].font = _font(18, True)
    sheet.row_dimensions[2].height = 32
    sheet['B3'] = _text(campaign_name)
    sheet['B3'].data_type = 's'
    sheet['B3'].font = _font(11, True)
    sheet['B4'] = scope
    sheet['B4'].font = _font(color=MUTED)
    sheet.merge_cells('E4:F4')
    sheet['E4'] = datetime.now()
    sheet['E4'].number_format = '"Édité le "dd/mm/yyyy" à "hh:mm'
    sheet['E4'].font = _font(color=MUTED)
    sheet['E4'].alignment = Alignment(horizontal='right')
    for column in range(2, 7):
        sheet.cell(5, column).border = Border(bottom=Side(style='medium', color=COBALT))
    sheet.row_dimensions[5].height = 6

    def heading(row, column, text):
        cell = sheet.cell(row, column, text)
        cell.font = _font(12, True)
        sheet.row_dimensions[row].height = 26

    def line(row, column, label, value, *, colors=None, link=None):
        name = sheet.cell(row, column, label)
        number = sheet.cell(row, column + 1, value)
        name.font = _font()
        number.font = _font(bold=True)
        number.alignment = Alignment(horizontal='right')
        for cell in (name, number):
            cell.border = Border(bottom=Side(style='thin', color=ROW_LINE))
        if colors:
            _chip(name, colors)
        if link:
            name.hyperlink = Hyperlink(ref=name.coordinate, location=link, display=label)
        sheet.row_dimensions[row].height = 22

    situations = Counter(row['situation'] for row in rows)
    heading(7, 2, 'Où en sont les dossiers')
    current = 8
    for key in URGENCY:
        if situations[key]:
            line(
                current,
                2,
                SUMMARY_LABELS[key],
                situations[key],
                colors=SITUATION_COLORS[key],
                link=f"'{DOSSIER_SHEET}'!A6",
            )
            current += 1
    line(current, 2, 'Total', len(rows))
    sheet.cell(current, 2).font = _font(bold=True)
    current += 1

    mails = Counter('Envoyé' if row['mail'].startswith('Envoyé') else row['mail'] for row in rows if row['mail'])
    pointable = [row for row in rows if row['dossier']['status'] == 'TERMINE']
    pointed = sum(bool(row['dossier'].get('pointed_at')) for row in pointable)
    heading(7, 5, 'Courriels et pointage')
    for offset, (label, value) in enumerate(
        (
            ('Mails prêts à envoyer', mails['Prêt']),
            ('Ouverts, à confirmer', mails['À confirmer']),
            ('Envoyés', mails['Envoyé']),
            ('Adresses à ajouter', mails['Sans adresse']),
            ('Dossiers pointés', pointed),
            ('Dossiers à pointer', len(pointable) - pointed),
        )
    ):
        line(8 + offset, 5, label, value)

    todo_row = max(current, 14) + 1
    heading(todo_row, 2, 'À faire')
    actions = [
        (
            situations['MANUAL'],
            _plural(situations['MANUAL'], 'dossier à vérifier manuellement', 'dossiers à vérifier manuellement'),
            'Vérifiez-les sur impots.gouv.fr, puis cochez « Vérifié » dans leur fiche.',
        ),
        (
            situations['RETRY'],
            _plural(
                situations['RETRY'],
                'dossier à relancer ou sans habilitation',
                'dossiers à relancer ou sans habilitation',
            ),
            'Relancez-les depuis l’onglet Collecte, ou vérifiez les comptes experts.',
        ),
        (
            situations['REVIEW'],
            _plural(situations['REVIEW'], 'lecture à vérifier', 'lectures à vérifier'),
            'Ouvrez l’avis et corrigez la lecture dans la fiche du dossier.',
        ),
        (
            situations['WAIT'],
            _plural(situations['WAIT'], 'dossier pas encore collecté', 'dossiers pas encore collectés'),
            'Lancez ou reprenez la collecte.',
        ),
        (
            mails['Prêt'],
            _plural(mails['Prêt'], 'mail prêt à envoyer', 'mails prêts à envoyer'),
            'Envoyez-les depuis l’onglet Courriels.',
        ),
        (
            mails['Sans adresse'],
            _plural(mails['Sans adresse'], 'adresse e-mail à ajouter', 'adresses e-mail à ajouter'),
            'Ajoutez-les dans l’onglet Courriels.',
        ),
    ]
    actions = [action for action in actions if action[0]] or [
        (0, 'Rien à faire pour l’instant', 'Tous les dossiers sont traités.')
    ]
    for offset, (_count, title, instruction) in enumerate(actions[:5]):
        row = todo_row + 1 + offset
        sheet.cell(row, 2, title).font = _font(bold=True)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        sheet.cell(row, 3, instruction).font = _font(color=MUTED)
        for column in range(2, 7):
            sheet.cell(row, column).border = Border(bottom=Side(style='thin', color=ROW_LINE))
        sheet.row_dimensions[row].height = 22
    last = todo_row + 1 + len(actions[:5])

    if cabinet:
        table_row = last + 1
        heading(table_row, 2, 'Par collaborateur')
        headers = [
            'Collaborateur',
            'Dossiers',
            'À traiter',
            'À régler',
            'Prélevé',
            'Rien à payer',
            'Sans avis',
            'Mails envoyés',
            'Pointés',
        ]
        for index, header in enumerate(headers, 2):
            cell = sheet.cell(table_row + 1, index, header)
            cell.fill = _fill(INK)
            cell.font = _font(bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal='left' if index == 2 else 'right', vertical='center')
        people = {}
        for row in rows:
            people.setdefault(row['dossier']['collaborator'] or 'À affecter', []).append(row)
        for offset, (name, items) in enumerate(sorted(people.items(), key=lambda item: item[0].casefold())):
            counts = Counter(item['situation'] for item in items)
            values = [
                name,
                len(items),
                counts['MANUAL'] + counts['RETRY'] + counts['REVIEW'] + counts['WAIT'],
                counts['PAY'],
                counts['DEBIT'],
                counts['NONE'],
                counts['ABSENT'],
                sum(item['mail'].startswith('Envoyé') for item in items),
                sum(bool(item['dossier'].get('pointed_at')) for item in items),
            ]
            for index, value in enumerate(values, 2):
                cell = sheet.cell(table_row + 2 + offset, index, value)
                cell.font = _font(bold=index == 2)
                cell.border = Border(bottom=Side(style='thin', color=ROW_LINE))
                if index > 2:
                    cell.alignment = Alignment(horizontal='right')
            sheet.row_dimensions[table_row + 2 + offset].height = 22
        last = table_row + 2 + len(people)

    years = sorted({row['year'] for row in rows})
    if len(years) > 1:
        table_row = last + 1
        heading(table_row, 2, 'Par année')
        for index, header in enumerate(['Année', 'Dossiers', 'Avec avis', 'À traiter', 'Sans avis'], 2):
            cell = sheet.cell(table_row + 1, index, header)
            cell.fill = _fill(INK)
            cell.font = _font(bold=True, color=WHITE)
        for offset, year in enumerate(years):
            items = [row for row in rows if row['year'] == year]
            counts = Counter(item['situation'] for item in items)
            values = [
                year,
                len(items),
                counts['PAY'] + counts['DEBIT'] + counts['NONE'],
                counts['MANUAL'] + counts['RETRY'] + counts['REVIEW'] + counts['WAIT'],
                counts['ABSENT'],
            ]
            for index, value in enumerate(values, 2):
                cell = sheet.cell(table_row + 2 + offset, index, value)
                cell.font = _font()
                cell.border = Border(bottom=Side(style='thin', color=ROW_LINE))
        last = table_row + 2 + len(years)

    sheet.print_area = f'A1:{"J" if cabinet else "G"}{last + 2}'
    sheet.page_setup.orientation = 'landscape' if cabinet else 'portrait'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1


def action_workbook(campaign_name, scope, records, notices, *, cabinet=False):
    """One workbook: « Synthèse », then « Dossiers » as in the application, then « Avis » to check each notice.

    records: one per dossier and year, with its classified notices, last year's explanation and its folder.
    notices: one per notice found, classified or not."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Synthèse'
    dossiers = workbook.create_sheet(DOSSIER_SHEET)
    avis = workbook.create_sheet('Avis')
    rows = dossier_rows(records)
    _summary(summary, campaign_name, scope, rows, cabinet)

    headers = [
        'Pointé le',
        'Dossier',
        'SIREN',
        'Collaborateur',
        'Année',
        'À payer',
        'Situation',
        'Échéance',
        'Compte prélevé',
        'Évolution sur un an',
        'Mail',
        'E-mail du client',
        'Avis',
        'Dossier client',
        'Capture du portail',
        'Information',
    ]
    values = []
    pdfs = []
    folders = []
    proofs = []
    for row in rows:
        dossier, classified = row['dossier'], row['notices']
        due = row['due']
        values.append(
            [
                _date(dossier.get('pointed_at') or '', with_time=True),
                _text(dossier['name']),
                _text(dossier['siren']),
                _text(dossier['collaborator']),
                row['year'],
                due / 100 if classified and due is not None else None,
                row['label'],
                _date(earliest_due_date(classified)),
                _text(debited_account(classified)),
                '\n'.join(row.get('evolution') or []),
                row['mail'],
                _text(dossier.get('email')),
                '',
                '',
                '',
                _information(row),
            ]
        )
        pdfs.append(classified[0]['path'] if len(classified) == 1 and Path(classified[0]['path']).is_file() else None)
        folders.append(row.get('folder') if row.get('folder') and Path(row['folder']).is_dir() else None)
        proof = dossier.get('proof_path') or ''
        proofs.append(
            proof if dossier['status'] in ('ABSENT', 'SANS_HABILITATION') and proof and Path(proof).is_file() else None
        )
    _sheet_heading(
        dossiers,
        'Dossiers',
        f'{campaign_name}  ·  {scope}  ·  {_plural(len(rows), "dossier", "dossiers")}'
        + (' / année' if len({row['year'] for row in rows}) > 1 else ''),
        '',
    )
    _table(
        dossiers,
        headers,
        values,
        [17, 30, 13, 16, 8, 14, 22, 12, 34, 46, 14, 30, 16, 16, 18, 52],
        'TableauDossiers',
        money=('À payer',),
        dates=('Échéance',),
        datetimes=('Pointé le',),
        texts=('SIREN', 'Compte prélevé'),
        wrap=('Évolution sur un an', 'Information'),
        links={
            'Avis': (pdfs, 'Ouvrir l’avis'),
            'Dossier client': (folders, 'Ouvrir le dossier'),
            'Capture du portail': (proofs, 'Ouvrir la capture'),
        },
    )
    situation_column = headers.index('Situation') + 1
    mail_column = headers.index('Mail') + 1
    for offset, row in enumerate(rows):
        _chip(dossiers.cell(7 + offset, situation_column), SITUATION_COLORS[row['situation']])
        mail = row['mail']
        if mail:
            _chip(dossiers.cell(7 + offset, mail_column), MAIL_COLORS.get(mail, ('DFF1F6', '1A6A80')))
        if len(row['notices']) > 1 and folders[offset]:
            _link(dossiers.cell(7 + offset, headers.index('Avis') + 1), folders[offset], f'{len(row["notices"])} avis')

    notice_headers = [
        'Dossier',
        'SIREN',
        'Année',
        'Type',
        'Référence',
        'Montant de l’avis',
        'À payer',
        'Mode de paiement',
        'Échéance',
        'Commune',
        'Lecture',
        'PDF',
    ]
    notice_values = []
    notice_links = []
    for record in sorted(
        notices, key=lambda item: (item['dossier']['name'].casefold(), item['year'], item['notice'].get('kind') or '')
    ):
        dossier, notice = record['dossier'], record['notice']
        complete = notice['status'] == 'TERMINE'
        notice_values.append(
            [
                _text(dossier['name']),
                _text(dossier['siren']),
                record['year'],
                KIND_LABELS.get(notice['kind'], _text(notice['kind'])),
                _text(notice['reference']),
                notice['tax_amount_cents'] / 100 if complete and notice['tax_amount_cents'] is not None else None,
                notice['amount_due_cents'] / 100 if complete and notice['amount_due_cents'] is not None else None,
                PAYMENT_LABELS.get(notice.get('payment_method'), 'Non indiqué sur l’avis') if complete else '',
                _date(notice['due_date']),
                _text(notice['commune']),
                review_label(notice) if complete else 'À relancer · ' + (_text(notice['error']) or 'avis non classé'),
                '',
            ]
        )
        notice_links.append(notice['path'] if complete else None)
    _sheet_heading(
        avis, 'Avis', f'{campaign_name}  ·  {scope}  ·  {_plural(len(notices), "avis trouvé", "avis trouvés")}', ''
    )
    _table(
        avis,
        notice_headers,
        notice_values,
        [30, 13, 8, 18, 20, 16, 14, 40, 12, 20, 30, 16],
        'TableauAvis',
        money=('Montant de l’avis', 'À payer'),
        dates=('Échéance',),
        texts=('SIREN', 'Référence'),
        wrap=('Lecture',),
        links={'PDF': (notice_links, 'Ouvrir le PDF')},
    )
    for offset, values_row in enumerate(notice_values):
        mode = values_row[notice_headers.index('Mode de paiement')]
        if isinstance(mode, str) and mode.startswith('À régler'):
            avis.cell(7 + offset, notice_headers.index('Mode de paiement') + 1).font = _font(bold=True, color='A85600')
        if values_row[notice_headers.index('Lecture')].startswith('À relancer'):
            avis.cell(7 + offset, notice_headers.index('Lecture') + 1).font = _font(bold=True, color='B01F48')
    for sheet, color in ((summary, COBALT), (dossiers, INK), (avis, MUTED)):
        sheet.sheet_properties.tabColor = color
    return workbook
