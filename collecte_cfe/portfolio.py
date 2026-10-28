import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from .domain import Dossier, safe_name

REQUIRED = ('Nom du dossier', 'SIREN', 'Collaborateur')
FIELD_MESSAGES = {
    'nom du dossier': 'nom du dossier manquant',
    'collaborateur': 'collaborateur manquant',
    'SIREN': 'SIREN manquant ou invalide (9 chiffres requis)',
    'e-mail': 'e-mail du client invalide',
}


@dataclass(frozen=True)
class PortfolioIssue:
    line: int
    level: str
    message: str


@dataclass
class PortfolioCheck:
    """Every dossier kept from a portfolio file, and every line ignored (error) or kept with a doubt (warning)."""

    dossiers: list[Dossier] = field(default_factory=list)
    issues: list[PortfolioIssue] = field(default_factory=list)
    invalid_rows: list[tuple[int, tuple[str, ...]]] = field(default_factory=list)

    @property
    def ignored(self):
        return [issue for issue in self.issues if issue.level == 'error']

    @property
    def warnings(self):
        return [issue for issue in self.issues if issue.level == 'warning']


def valid_client_email(value):
    return bool(re.fullmatch(r'[^@\s;,<>]+@[^@\s;,<>]+\.[^@\s;,<>]+', str(value or '').strip()))


def normalized_siren(value):
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        return str(int(value)).zfill(9) if float(value).is_integer() else ''
    return re.sub(r'\s', '', str(value or ''))


def siren_check_digit_is_valid(siren):
    total = 0
    for index, character in enumerate(reversed(siren)):
        digit = int(character) * (2 if index % 2 else 1)
        total += digit - 9 if digit > 9 else digit
    return total % 10 == 0


def looks_like_headerless_portfolio(row):
    # Without headers: name, SIREN, collaborator, and optionally the client's e-mail as a fourth column.
    if len(row) < 3 or any(v is not None and str(v).strip() for v in row[4:]):
        return False
    fourth = str(row[3] or '').strip() if len(row) > 3 else ''
    if fourth and '@' not in fourth:
        return False
    return bool(
        str(row[0] or '').strip() and re.fullmatch(r'[0-9]{9}', normalized_siren(row[1])) and str(row[2] or '').strip()
    )


def _read_rows(path: Path):
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(book.active.iter_rows(values_only=True))
        finally:
            book.close()
    if path.suffix.lower() == '.csv':
        raw = path.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1252')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            raise ValueError('CSV illisible : utilisez des virgules ou des points-virgules.') from None
        return list(csv.reader(text.splitlines(), dialect))
    raise ValueError('Choisissez un portefeuille CSV ou XLSX.')


def check_portfolio(path: Path | str) -> PortfolioCheck:
    """Read a portfolio and report every problem by line instead of stopping at the first one."""
    rows = _read_rows(Path(path))
    if not rows:
        raise ValueError('Le portefeuille est vide.')
    aliases = {
        'nom du dossier': 'Nom du dossier',
        'siren': 'SIREN',
        'collaborateur': 'Collaborateur',
        'nom du collab': 'Collaborateur',
        'nom du collaborateur': 'Collaborateur',
        'e-mail du client': 'E-mail du client',
        'email du client': 'E-mail du client',
        'e-mail': 'E-mail du client',
        'email': 'E-mail du client',
        'mail client': 'E-mail du client',
        'mail': 'E-mail du client',
        'courriel': 'E-mail du client',
        'e-mail client': 'E-mail du client',
        'email client': 'E-mail du client',
        'adresse mail': 'E-mail du client',
        'adresse e-mail': 'E-mail du client',
        'adresse email': 'E-mail du client',
        'courriel du client': 'E-mail du client',
        'mail du client': 'E-mail du client',
    }
    headers = [aliases.get(str(v or '').strip().casefold(), str(v or '').strip()) for v in rows[0]]
    missing = [h for h in REQUIRED if headers.count(h) != 1]
    if missing:
        if all(headers.count(h) == 0 for h in REQUIRED) and looks_like_headerless_portfolio(rows[0]):
            indexes = [0, 1, 2]
            data_rows = rows
            first_line = 1
            headers = ['Nom du dossier', 'SIREN', 'Collaborateur', 'E-mail du client']
        else:
            raise ValueError('Colonnes obligatoires absentes ou répétées : ' + ', '.join(missing))
    else:
        indexes = [headers.index(h) for h in REQUIRED]
        data_rows = rows[1:]
        first_line = 2
    email_index = headers.index('E-mail du client') if headers.count('E-mail du client') == 1 else None
    if headers.count('E-mail du client') > 1:
        raise ValueError('La colonne e-mail du client apparaît plusieurs fois.')
    check = PortfolioCheck()
    siren_lines = {}
    destination_lines = {}
    collaborators = {}
    for line, row in enumerate(data_rows, first_line):
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        values = [row[i] if i < len(row) else None for i in indexes]
        name, siren, collaborator = values
        siren = normalized_siren(siren)
        name = str(name or '').strip()
        collaborator = str(collaborator or '').strip()
        email = str(row[email_index] or '').strip() if email_index is not None and email_index < len(row) else ''
        invalid = []
        if not name:
            invalid.append('nom du dossier')
        if not collaborator:
            invalid.append('collaborateur')
        if not re.fullmatch(r'[0-9]{9}', siren):
            invalid.append('SIREN')
        if email and not valid_client_email(email):
            invalid.append('e-mail')
        if invalid:
            check.invalid_rows.append((line, tuple(invalid)))
            check.issues.append(PortfolioIssue(line, 'error', ' ; '.join(FIELD_MESSAGES[field] for field in invalid)))
            continue
        if siren in siren_lines:
            check.issues.append(
                PortfolioIssue(line, 'error', f'SIREN en doublon ({siren}), déjà présent ligne {siren_lines[siren]}')
            )
            continue
        destination = (safe_name(collaborator).casefold(), safe_name(name).casefold())
        if destination in destination_lines:
            check.issues.append(
                PortfolioIssue(
                    line, 'error', f'même chemin de classement que la ligne {destination_lines[destination]}'
                )
            )
            continue
        # Distinct collaborator strings must not silently merge their action workbooks.
        other = collaborators.get(destination[0])
        if other is not None and other != collaborator:
            check.issues.append(
                PortfolioIssue(
                    line, 'error', f'collision de classement entre collaborateurs (« {other} » et « {collaborator} »)'
                )
            )
            continue
        siren_lines[siren] = line
        destination_lines[destination] = line
        collaborators[destination[0]] = collaborator
        if not siren_check_digit_is_valid(siren):
            check.issues.append(
                PortfolioIssue(line, 'warning', f'SIREN {siren} à vérifier : la clé de contrôle ne correspond pas')
            )
        check.dossiers.append(Dossier(name, siren, collaborator, email))
    return check


def import_portfolio(path: Path | str) -> list[Dossier]:
    """Strict import: any line that would be ignored stops the import. Check-digit warnings do not."""
    check = check_portfolio(path)
    if check.invalid_rows:
        lines = [line for line, _fields in check.invalid_rows]
        prefix = ('Ligne ' if len(lines) == 1 else 'Lignes ') + ', '.join(str(line) for line in lines[:12])
        if len(lines) > 12:
            prefix += f' et {len(lines) - 12} autre(s)'
        field_sets = {fields for _line, fields in check.invalid_rows}
        if field_sets == {('SIREN',)}:
            detail = 'SIREN manquant ou invalide (9 chiffres requis).'
        elif field_sets == {('e-mail',)}:
            detail = 'e-mail du client invalide.'
        else:
            detail = 'nom du dossier, collaborateur et SIREN à 9 chiffres obligatoires ; vérifiez aussi les e-mails.'
        raise ValueError(f'{prefix} : {detail}')
    if check.ignored:
        issue = check.ignored[0]
        raise ValueError(f'Ligne {issue.line} : {issue.message}.')
    if not check.dossiers:
        raise ValueError('Le portefeuille ne contient aucun dossier.')
    return check.dossiers
