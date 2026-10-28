import os

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtCore import Qt

from collecte_cfe.dashboard import MainWindow
from collecte_cfe.domain import Dossier, Notice
from collecte_cfe.notice_extraction import NoticeExtraction
from collecte_cfe.store import Store
from tests.test_dashboard import MemorySettings


def test_mail_page_opens_only_selected_ready_client_and_saves_address(tmp_path, qapp):
    store = Store(tmp_path / 'db.sqlite')
    dossiers = [Dossier('Alpha', '000000001', 'Léa', 'alpha@example.fr'), Dossier('Beta', '000000003', 'Léa')]
    batch = store.create_batch([2025], tmp_path / 'sortie', dossiers, 1)
    campaign = store.batch_campaigns(batch)[0]['id']
    for dossier in dossiers:
        row = store.register_notice(
            campaign,
            dossier.siren,
            Notice(2025, '123' + dossier.siren, '00021', '33', 'BORDEAUX', 'Rôle général', '15/12/2025'),
        )
        store.notice_result(row['id'], 'TERMINE')
        store.notice_extraction_result(row['id'], NoticeExtraction(10000, 10000, 'A_REGLER', '', False, 'EXTRAIT'))
        store.dossier_result(campaign, dossier.siren, 'TERMINE')
    w = MainWindow(store, settings=MemorySettings())
    w.navigate(3)
    assert w.title.text() == 'Campagne 1' and w.campaign_tab_buttons[3].isChecked()
    assert w.mail_table.rowCount() == 2
    assert not w.mail_table.item(1, 0).flags() & Qt.ItemFlag.ItemIsEnabled
    w.mail_table.item(1, 3).setText('beta@example.fr')
    assert store.dossiers(campaign)[1]['email'] == 'beta@example.fr'
    errors = []
    w.problem = errors.append
    w.mail_table.item(1, 3).setText('adresse invalide')
    assert store.dossiers(campaign)[1]['email'] == 'beta@example.fr'
    assert w.mail_table.item(1, 3).text() == 'beta@example.fr'
    assert errors and 'invalide' in errors[0]
    w.mail_table.item(0, 0).setCheckState(Qt.CheckState.Checked)
    drafts = []
    w.draft_opener = lambda to, subject, body, files: drafts.append((to, subject, body, files))
    w.open_selected_drafts()
    assert [to for to, *_rest in drafts] == ['alpha@example.fr']
    w.close()


def test_a_mail_opens_in_outlook_with_the_notice_pdfs_attached(tmp_path, qapp):
    from pathlib import Path

    store, batch, campaign = mail_store(tmp_path)
    for notice in store.notices(campaign):
        path = Path(notice['path'])
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b'%PDF-1.4 test')
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    w.navigate(3)
    drafts = []
    w.draft_opener = lambda to, subject, body, files: drafts.append((to, subject, body, files))
    w.mail_table.item(0, 0).setCheckState(Qt.CheckState.Checked)
    assert w.mail_open_button.text() == 'Envoyer 1 mail par Outlook'
    w.mail_open_button.click()
    ((to, subject, body, files),) = drafts
    assert to == 'alpha@example.fr' and subject == 'Avis CFE 2025 — Alpha'
    assert 'Vous trouverez votre avis en pièce jointe.' in body
    assert [path.name for path in files] == [Path(store.notices(campaign)[0]['path']).name]
    assert w.mail_filter_buttons['OPENED'].text() == 'À confirmer 1' and w.mail_filter_buttons['OPENED'].isChecked()
    w.mail_filter_buttons['READY'].click()
    w.mail_table.selectRow(next(r for r in range(w.mail_table.rowCount()) if w.mail_table.item(r, 1).text() == 'Gamma'))
    assert 'Gamma_CFE_2025_SOLDE.pdf' in w.mail_preview_files.text() and 'attachment.svg' in w.mail_preview_files.text()
    w.close()


def test_without_a_compatible_mail_application_the_mail_opens_without_attachments(tmp_path, qapp, monkeypatch):
    from collecte_cfe.outlook_draft import DraftUnavailable

    store, batch, campaign = mail_store(tmp_path)
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    w.navigate(3)

    def unavailable(*_args):
        raise DraftUnavailable('Aucune messagerie compatible n’est installée.')

    w.draft_opener = unavailable
    opened = []
    w.mail_opener = lambda url: opened.append(url.toString()) or True
    problems = []
    monkeypatch.setattr(w, 'problem', problems.append)
    monkeypatch.setattr('collecte_cfe.ui_mail.QDesktopServices.openUrl', lambda url: True)
    w.mail_table.item(0, 0).setCheckState(Qt.CheckState.Checked)
    w.open_selected_drafts()
    assert len(opened) == 1 and opened[0].startswith('mailto:alpha@example.fr?')
    assert problems and problems[0].startswith('Aucune messagerie compatible') and 'sans pièce jointe' in problems[0]
    w.close()


def mail_store(tmp_path):
    store = Store(tmp_path / 'db.sqlite')
    dossiers = [
        Dossier('Alpha', '000000001', 'Léa', 'alpha@example.fr'),
        Dossier('Beta', '000000003', 'Léa'),
        Dossier('Gamma', '000000004', 'Paul', 'gamma@example.fr'),
        Dossier('Delta', '000000005', 'Paul', 'delta@example.fr'),
    ]
    batch = store.create_batch([2025], tmp_path / 'sortie', dossiers, 1, 'CFE test')
    campaign = store.batch_campaigns(batch)[0]['id']
    readings = {
        '000000001': (128400, 'A_REGLER'),
        '000000003': (10000, 'A_REGLER'),
        '000000004': (23100, 'PRELEVEMENT_ECHEANCE'),
    }
    for siren, (due, method) in readings.items():
        row = store.register_notice(
            campaign, siren, Notice(2025, 'R' + siren, '00021', '33', 'BORDEAUX', 'Rôle général', '15/12/2025')
        )
        store.notice_result(row['id'], 'TERMINE')
        store.notice_extraction_result(row['id'], NoticeExtraction(due, due, method, '', False, 'EXTRAIT'))
        store.dossier_result(campaign, siren, 'TERMINE')
    return store, batch, campaign


def test_the_mail_page_sorts_clients_and_says_what_each_message_says(tmp_path, qapp):
    store, batch, campaign = mail_store(tmp_path)
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    w.campaign_tab_buttons[3].click()
    assert [chip.text() for chip in w.mail_filter_buttons.values()] == [
        'Tous 4',
        'Prêts 2',
        'À confirmer 0',
        'Envoyés 0',
        'Sans adresse 1',
        'En attente 1',
    ]
    assert w.mail_filter_buttons['ALL'].isChecked() and w.campaign_tab_buttons[3].badge.text() == '2 prêts'
    rows = [
        (
            w.mail_table.item(row, 1).text(),
            w.mail_table.item(row, 4).text(),
            w.mail_table.item(row, 5).text(),
            w.mail_table.cellWidget(row, 6).action.text(),
        )
        for row in range(w.mail_table.rowCount())
        if not w.mail_table.isRowHidden(row)
    ]
    assert sorted(rows) == [
        ('Alpha', '1 284,00 € à régler', 'Prêt', 'Envoyer'),
        ('Beta', '100,00 € à régler', 'Sans adresse', 'Ajouter l’adresse'),
        ('Delta', 'Collecte à terminer', 'En attente', 'Voir le dossier'),
        ('Gamma', 'Prélevé à l’échéance', 'Prêt', 'Envoyer'),
    ]
    assert w.mail_preview_to.text() == 'À : alpha@example.fr', 'the first mail is shown before a client is chosen'
    w.mail_filter_buttons['READY'].click()
    visible = [row for row in range(w.mail_table.rowCount()) if not w.mail_table.isRowHidden(row)]
    assert [w.mail_table.item(row, 1).text() for row in visible] == ['Alpha', 'Gamma']
    w.mail_filter_buttons['NO_EMAIL'].click()
    assert [
        w.mail_table.item(row, 1).text() for row in range(w.mail_table.rowCount()) if not w.mail_table.isRowHidden(row)
    ] == ['Beta']
    w.mail_table.selectRow(0)
    assert w.mail_preview_to.text() == 'À : alpha@example.fr' and w.mail_preview.text().startswith('Bonjour,')
    w.close()


def test_an_opened_mail_waits_to_be_marked_sent_and_the_mark_is_kept(tmp_path, qapp):
    store, batch, campaign = mail_store(tmp_path)
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    w.navigate(3)
    assert not w.mail_open_button.isEnabled() and w.mail_open_button.text() == 'Envoyer les mails par Outlook'
    w.mail_select_toggle.click()
    assert (
        w.mail_open_button.text() == 'Envoyer 2 mails par Outlook'
        and w.mail_select_toggle.text() == 'Tout désélectionner'
    )
    drafts = []
    w.draft_opener = lambda *args: drafts.append(args)
    w.mail_open_button.click()
    assert len(drafts) == 2
    assert w.mail_filter_buttons['OPENED'].isChecked() and w.mail_filter_buttons['OPENED'].text() == 'À confirmer 2'
    rows = [row for row in range(w.mail_table.rowCount()) if not w.mail_table.isRowHidden(row)]
    assert [w.mail_table.cellWidget(row, 6).action.text() for row in rows] == ['Marquer envoyé'] * 2
    w.mail_table.cellWidget(rows[0], 6).action.click()
    assert (
        w.mail_filter_buttons['SENT'].text() == 'Envoyés 1'
        and w.mail_filter_buttons['OPENED'].text() == 'À confirmer 1'
    )
    w.mail_select_toggle.click()
    assert w.mail_open_button.text() == 'Marquer 1 mail comme envoyé'
    w.mail_open_button.click()
    assert w.mail_filter_buttons['SENT'].text() == 'Envoyés 2'
    assert all(
        w.mail_table.item(row, 5).text().startswith('Envoyé le ')
        for row in range(w.mail_table.rowCount())
        if not w.mail_table.isRowHidden(row)
    )
    w.close()
    reopened = MainWindow(store, settings=MemorySettings())
    reopened.open_campaign(batch)
    reopened.navigate(3)
    reopened.mail_filter_buttons['SENT'].click()
    rows = [row for row in range(reopened.mail_table.rowCount()) if not reopened.mail_table.isRowHidden(row)]
    assert all(reopened.mail_table.item(row, 5).text().startswith('Envoyé le ') for row in rows) and len(rows) == 2
    reopened.mail_table.item(rows[0], 0).setCheckState(Qt.CheckState.Checked)
    assert (
        reopened.mail_open_button.text() == 'Marquer 1 mail non envoyé'
        and reopened.mail_chosen.text() == '1 client choisi'
    )
    reopened.mail_open_button.click()
    assert (
        reopened.mail_filter_buttons['SENT'].text() == 'Envoyés 1'
        and reopened.mail_filter_buttons['READY'].text() == 'Prêts 0'
    )
    assert reopened.mail_filter_buttons['OPENED'].text() == 'À confirmer 1', (
        'a mail put back stays opened, to confirm again'
    )
    reopened.close()


def test_the_mail_summary_names_what_the_client_has_to_do():
    from collecte_cfe.ui_mail import mail_summary

    notice = {'status': 'TERMINE', 'payment_method': 'A_REGLER', 'amount_due_cents': 24300}
    assert mail_summary([notice, {**notice, 'amount_due_cents': 48600}]) == '729,00 € à régler'
    assert mail_summary([{**notice, 'payment_method': 'PRELEVEMENT_MENSUEL'}]) == 'Prélèvement mensuel'
    assert (
        mail_summary([{**notice, 'payment_method': 'AUCUN_PAIEMENT', 'refund_cents': 1500}])
        == 'Rien à payer, remboursement'
    )
    assert mail_summary([]) == ''


def test_the_dossier_sheet_sends_then_marks_the_mail_sent(tmp_path, qapp, monkeypatch):
    store, batch, campaign = mail_store(tmp_path)
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    monkeypatch.setattr(w.toast, 'announce', lambda text, *_args, **_kwargs: None)
    drafts = []
    w.draft_opener = lambda *args: drafts.append(args)
    row = next(r for r in range(w.table.rowCount()) if w.table.item(r, 0).text() == 'Alpha')
    w.table.selectRow(row)
    assert w.result_mail_button.text() == 'Envoyer le mail'
    w.result_mail_button.click()
    assert len(drafts) == 1 and w.result_mail_button.text() == 'Marquer comme envoyé'
    assert 'à confirmer une fois parti' in w.result_email_label.text()
    w.result_mail_button.click()
    assert w.result_mail_button.text() == 'Marquer non envoyé' and 'Mail envoyé le ' in w.result_email_label.text()
    assert store.dossiers(campaign)[0]['mail_sent_at']
    w.close()


def test_the_follow_up_sheet_shows_when_the_mail_was_sent(tmp_path):
    from datetime import datetime

    from openpyxl import load_workbook

    from collecte_cfe.documents import export_batch_actions

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch(
        [2025],
        tmp_path / 'out',
        [Dossier('Alpha', '000000001', 'Léa', 'alpha@example.fr'), Dossier('Beta', '000000002', 'Léa')],
        1,
    )
    campaign = store.batch_campaigns(batch)[0]['id']
    for siren in ('000000001', '000000002'):
        store.dossier_result(campaign, siren, 'TERMINE')
    store.mark_mail_sent(batch, '000000001', True)
    sheet = load_workbook(export_batch_actions(store, batch)[0])['Dossiers']
    headers = [cell.value for cell in sheet[6]]
    mails = {
        sheet.cell(row, 2).value: sheet.cell(row, headers.index('Mail') + 1).value
        for row in range(7, sheet.max_row + 1)
    }
    assert mails['Alpha'] == 'Envoyé ' + datetime.now().strftime('%d/%m') and not mails['Beta']


def test_outlook_opens_the_mail_without_freezing_the_window(tmp_path, qapp, monkeypatch):
    """Outlook can take seconds to start: the mail opens in the background and the window keeps answering."""
    import threading
    import time
    from pathlib import Path

    monkeypatch.setattr('collecte_cfe.ui_mail.DRAFTS_IN_BACKGROUND', True)
    store, batch, campaign = mail_store(tmp_path)
    for notice in store.notices(campaign):
        path = Path(notice['path'])
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b'%PDF-1.4 test')
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    w.navigate(3)
    release = threading.Event()
    threads = []

    def slow_outlook(*_args):
        threads.append(threading.current_thread())
        release.wait(5)

    w.draft_opener = slow_outlook
    w.mail_table.item(0, 0).setCheckState(Qt.CheckState.Checked)
    started = time.perf_counter()
    w.mail_open_button.click()
    assert time.perf_counter() - started < 0.5, 'the click returns before Outlook answers'
    assert w.mail_job is not None and not w.mail_open_button.isEnabled()
    assert w.toast.message.text() == 'Ouverture de 1 mail dans Outlook…'
    w.mail_open_button.click()
    assert len(threads) <= 1, 'a second click does not open the mail twice'
    release.set()
    deadline = time.monotonic() + 5
    while w.mail_job is not None and time.monotonic() < deadline:
        qapp.processEvents()
        time.sleep(0.02)
    assert threads and threads[0] is not threading.main_thread()
    assert w.mail_job is None and store.dossiers(campaign)[0]['mail_opened_at']
    assert 'ouvert dans Outlook' in w.toast.message.text()
    w.close()


def test_an_opened_or_sent_mail_can_be_opened_again(tmp_path, qapp):
    store, batch, campaign = mail_store(tmp_path)
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    w.navigate(3)
    drafts = []
    w.draft_opener = lambda to, *_rest: drafts.append(to)
    w.mail_select_toggle.click()
    w.mail_open_button.click()
    assert drafts == ['alpha@example.fr', 'gamma@example.fr']
    rows = [row for row in range(w.mail_table.rowCount()) if not w.mail_table.isRowHidden(row)]
    holder = w.mail_table.cellWidget(rows[0], 6)
    assert holder.action.text() == 'Marquer envoyé' and holder.resend.text() == 'Renvoyer'
    holder.resend.click()
    assert drafts[-1] == 'alpha@example.fr' and w.mail_filter_buttons['OPENED'].text() == 'À confirmer 2'
    w.mail_table.cellWidget(rows[0], 6).action.click()
    w.mail_filter_buttons['SENT'].click()
    sent = [row for row in range(w.mail_table.rowCount()) if not w.mail_table.isRowHidden(row)]
    assert (
        w.mail_table.cellWidget(sent[0], 6).action.text() == 'Renvoyer'
        and w.mail_table.cellWidget(sent[0], 6).resend is None
    )
    w.mail_table.cellWidget(sent[0], 6).action.click()
    assert drafts[-1] == 'alpha@example.fr' and len(drafts) == 4
    assert not store.dossiers(campaign)[0]['mail_sent_at'], 'a mail sent again waits to be confirmed again'
    assert (
        w.mail_filter_buttons['OPENED'].text() == 'À confirmer 2'
        and w.mail_filter_buttons['SENT'].text() == 'Envoyés 0'
    )
    w.close()


def test_the_cabinet_rewrites_the_mail_and_every_draft_follows(tmp_path, qapp, monkeypatch):
    from PySide6.QtWidgets import QDialog

    from collecte_cfe import ui_mail
    from collecte_cfe.client_mail import MailTemplate

    store, batch, campaign = mail_store(tmp_path)
    w = MainWindow(store, settings=MemorySettings())
    w.open_campaign(batch)
    w.navigate(3)

    def write(dialog):
        assert (
            dialog.sample is not None
            and 'Alpha' in dialog.windowTitle() + dialog.preview.text() + dialog.sample['dossier']['name']
        )
        dialog.subject.setText('Votre CFE {annees}')
        dialog.body.setPlainText('Bonjour {client},\n\n{avis}\n\nÀ bientôt,')
        assert dialog.save_button.isEnabled() and 'Bonjour Alpha,' in dialog.preview.text()
        dialog.body.setPlainText('Bonjour,')
        assert not dialog.save_button.isEnabled() and '{avis}' in dialog.error.text()
        dialog.body.setPlainText('Bonjour {client},\n\n{avis}\n\nÀ bientôt,')
        dialog.save()
        return dialog.result()

    monkeypatch.setattr(ui_mail.MailTemplateDialog, 'exec', write)
    assert w.edit_mail_template()
    assert store.mail_template() == MailTemplate('Votre CFE {annees}', 'Bonjour {client},\n\n{avis}\n\nÀ bientôt,')
    drafts = []
    w.draft_opener = lambda to, subject, body, files: drafts.append((subject, body))
    w.mail_table.item(0, 0).setCheckState(Qt.CheckState.Checked)
    w.mail_open_button.click()
    ((subject, body),) = drafts
    assert subject == 'Votre CFE 2025' and body.startswith('Bonjour Alpha,') and '1 284,00 €' in body

    def back_to_default(dialog):
        dialog.reset()
        dialog.save()
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(ui_mail.MailTemplateDialog, 'exec', back_to_default)
    w.edit_mail_template()
    assert store.mail_template() == MailTemplate()
    w.close()
