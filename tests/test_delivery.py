import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from collecte_cfe.documents import classify_pdf, export_actions, proof_pdf
from collecte_cfe.domain import Dossier, Notice
from collecte_cfe.store import Store


def sample_pdf(path):
    c = canvas.Canvas(str(path))
    c.drawString(50, 700, 'Document fiscal original')
    c.showPage()
    c.save()


def test_resume_and_notice_numbering_survive_database_reopen(tmp_path):
    d = Dossier('Exemple', '000000001', 'Léa')
    store = Store(tmp_path / 'db.sqlite')
    campaign = store.create_campaign(2025, tmp_path / 'sortie', [d])
    n = Notice(2025, '123', '00021', '33', 'BORDEAUX', 'Rôle général n°092', '15/12/2025')
    row = store.register_notice(campaign, d.siren, n)
    assert Path(row['path']).name == 'Exemple_CFE_2025_SOLDE.pdf'
    source = tmp_path / 'resume.pdf'
    sample_pdf(source)
    digest = classify_pdf(source, Path(row['path']), d, n)
    store.notice_result(row['id'], 'TERMINE', '', digest)
    store = Store(tmp_path / 'db.sqlite')
    assert store.register_notice(campaign, d.siren, n)['status'] == 'TERMINE'
    other = Notice(2025, '124', '00022', '33', 'BORDEAUX', 'Rôle général n°092', '15/12/2025')
    assert Path(store.register_notice(campaign, d.siren, other)['path']).name == 'Exemple_CFE_2025_SOLDE_02.pdf'
    acompte = Notice(2025, '125', '00022', '33', 'BORDEAUX', 'Rôle d’acompte n°002', '15/06/2025')
    assert Path(store.register_notice(campaign, d.siren, acompte)['path']).name == 'Exemple_CFE_2025_ACOMPTE.pdf'


def test_classification_preserves_original_and_removes_temp_only_after_success(tmp_path):
    source = tmp_path / 'temp.pdf'
    sample_pdf(source)
    d = Dossier('Exemple', '000000001', 'Léa')
    n = Notice(2025, '123', '00021', '33', 'BORDEAUX', 'Rôle général n°092', '')
    target = tmp_path / 'Léa' / 'Exemple' / 'Exemple_CFE_2025_SOLDE.pdf'
    digest = classify_pdf(source, target, d, n)
    assert not source.exists()
    r = PdfReader(target)
    assert len(r.pages) == 1, 'the notice is filed without an added first page'
    assert 'Document fiscal original' in r.pages[0].extract_text()
    assert r.metadata['/CollecteNoticeKey'] == n.key and r.metadata['/CollecteCoverPage'] == 'NON'
    assert digest == hashlib.sha256(target.read_bytes()).hexdigest()


def test_html_download_cannot_be_classified(tmp_path):
    source = tmp_path / 'temp.pdf'
    source.write_text('<html>Connexion</html>')
    target = tmp_path / 'avis.pdf'
    with pytest.raises(ValueError, match='PDF'):
        classify_pdf(
            source, target, Dossier('X', '123456789', 'A'), Notice(2025, '123', '1', '33', 'X', 'Rôle général', '')
        )
    assert source.exists() and not target.exists()


def test_proof_and_actions_have_no_financial_total(tmp_path):
    d = Dossier('Exemple', '000000001', 'Léa')
    s = Store(tmp_path / 'db.sqlite')
    campaign = s.create_campaign(2025, tmp_path / 'sortie', [d])
    proof = s.register_proof(campaign, d.siren)
    digest = proof_pdf(proof, d, 2025)
    s.proof_result(campaign, d.siren, digest)
    s.dossier_result(campaign, d.siren, 'ABSENT', 'Aucun avis disponible pour 2025')
    assert '2025' in PdfReader(proof).pages[0].extract_text()
    paths = export_actions(s, campaign)
    assert len(paths) == 1
    w = load_workbook(paths[0])
    rows = list(w['Dossiers'].values)
    assert rows[6][1] == 'Exemple' and rows[6][2] == '000000001' and rows[6][6] == 'Aucun avis'
    assert not any('€' in str(cell.number_format) for row in w['Synthèse'].iter_rows() for cell in row), (
        'no total in euros'
    )


def test_existing_classified_notice_is_enriched_and_exported_without_total(tmp_path):
    d = Dossier('Exemple', '000000001', 'Léa')
    s = Store(tmp_path / 'db.sqlite')
    campaign = s.create_campaign(2025, tmp_path / 'sortie', [d])
    notice = Notice(2025, '123', '00021', '33', 'BORDEAUX', 'Rôle général', '15/12/2025')
    row = s.register_notice(campaign, d.siren, notice)
    source = tmp_path / 'temp.pdf'
    pdf = canvas.Canvas(str(source))
    for line, y in [
        ('MONTANT DE VOTRE IMPOT 1 234,00 EUR', 730),
        ('MONTANT A PAYER 934,00 EUR', 710),
        ("Vous avez choisi le PRELEVEMENT A L'ECHEANCE.", 690),
        ('sur le compte bancaire FR00 0000 0000 0000 XXXXXXXX 000.', 670),
    ]:
        pdf.drawString(50, y, line)
    pdf.showPage()
    pdf.save()
    classify_pdf(source, Path(row['path']), d, notice)
    s.notice_result(row['id'], 'TERMINE', '', hashlib.sha256(Path(row['path']).read_bytes()).hexdigest())

    workbook = load_workbook(export_actions(s, campaign)[0], data_only=True)
    sheet = workbook['Avis']
    headers = [cell.value for cell in sheet[6]]
    values = dict(zip(headers, [cell.value for cell in sheet[7]]))
    assert values['Montant de l’avis'] == 1234
    assert values['À payer'] == 934
    assert sheet.cell(7, headers.index('Montant de l’avis') + 1).number_format != 'General'
    assert values['Mode de paiement'] == 'Prélèvement à l’échéance'
    assert values['Lecture'] == 'Partielle · compte masqué'
    dossier_headers = [cell.value for cell in workbook['Dossiers'][6]]
    assert (
        workbook['Dossiers'].cell(7, dossier_headers.index('Compte prélevé') + 1).value
        == 'FR00 0000 0000 0000 XXXXXXXX 000'
    )
    assert sheet.cell(7, headers.index('PDF') + 1).hyperlink is not None
    saved = Store(tmp_path / 'db.sqlite').notices(campaign, d.siren)[0]
    assert saved['tax_amount_cents'] == 123400
    assert saved['amount_due_cents'] == 93400
    assert saved['bank_account_masked'] == 1


def test_modified_classified_notice_blocks_export_without_mutation(tmp_path):
    d = Dossier('Exemple', '000000001', 'Léa')
    s = Store(tmp_path / 'db.sqlite')
    campaign = s.create_campaign(2025, tmp_path / 'sortie', [d])
    # Create an existing workbook before adding the notice; a failed export must
    # preserve its bytes exactly.
    target = export_actions(s, campaign)[0]
    notice = Notice(2025, '123', '00021', '33', 'BORDEAUX', 'Rôle général', '15/12/2025')
    row = s.register_notice(campaign, d.siren, notice)
    source = tmp_path / 'avis-original.pdf'
    sample_pdf(source)
    digest = classify_pdf(source, Path(row['path']), d, notice)
    s.notice_result(row['id'], 'TERMINE', '', digest)
    Path(row['path']).write_bytes(Path(row['path']).read_bytes() + b'alteration')
    altered_notice = Path(row['path']).read_bytes()
    workbook_before = target.read_bytes()
    notices_before = s.notices(campaign)
    dossiers_before = s.dossiers(campaign)

    with pytest.raises(ValueError, match='000000001') as error:
        export_actions(s, campaign)

    assert 'relancez la collecte' in str(error.value).lower()
    assert s.notices(campaign) == notices_before
    assert s.dossiers(campaign) == dossiers_before
    assert Path(row['path']).read_bytes() == altered_notice
    assert target.read_bytes() == workbook_before


def test_modified_absence_proof_blocks_batch_export_without_mutation(tmp_path):
    from collecte_cfe.documents import export_batch_actions

    d = Dossier('Exemple', '000000001', 'Léa')
    s = Store(tmp_path / 'db.sqlite')
    batch = s.create_batch([2025], tmp_path / 'sortie', [d], 1)
    campaign = s.batch_campaigns(batch)[0]['id']
    proof = s.register_proof(campaign, d.siren)
    digest = proof_pdf(proof, d, 2025)
    s.proof_result(campaign, d.siren, digest)
    s.dossier_result(campaign, d.siren, 'ABSENT', 'Aucun avis disponible pour 2025')
    target = export_batch_actions(s, batch)[0]
    proof.write_bytes(proof.read_bytes() + b'alteration')
    altered_proof = proof.read_bytes()
    workbook_before = target.read_bytes()
    notices_before = s.notices(campaign)
    dossiers_before = s.dossiers(campaign)

    with pytest.raises(ValueError, match='000000001') as error:
        export_batch_actions(s, batch)

    assert 'relancez la collecte' in str(error.value).lower()
    assert s.notices(campaign) == notices_before
    assert s.dossiers(campaign) == dossiers_before
    assert proof.read_bytes() == altered_proof
    assert target.read_bytes() == workbook_before


@pytest.mark.parametrize('digest', ['', 'f' * 64], ids=['empreinte-absente', 'fichier-absent'])
def test_completed_notice_without_verified_file_blocks_export_before_writing(tmp_path, digest):
    d = Dossier('Exemple', '000000001', 'Léa')
    s = Store(tmp_path / 'db.sqlite')
    output = tmp_path / 'sortie'
    campaign = s.create_campaign(2025, output, [d])
    notice = Notice(2025, '123', '00021', '33', 'BORDEAUX', 'Rôle général', '15/12/2025')
    row = s.register_notice(campaign, d.siren, notice)
    s.notice_result(row['id'], 'TERMINE', '', digest)
    notices_before = s.notices(campaign)
    dossiers_before = s.dossiers(campaign)

    with pytest.raises(ValueError, match='000000001'):
        export_actions(s, campaign)

    assert s.notices(campaign) == notices_before
    assert s.dossiers(campaign) == dossiers_before
    assert not Path(row['path']).exists()
    assert not output.exists()


def test_corruption_in_later_campaign_precedes_earlier_notice_enrichment(tmp_path):
    from collecte_cfe.documents import export_batch_actions

    d = Dossier('Exemple', '000000001', 'Léa')
    s = Store(tmp_path / 'db.sqlite')
    batch = s.create_batch([2024, 2025], tmp_path / 'sortie', [d], 1)
    first, second = s.batch_campaigns(batch)
    first_notice = Notice(2024, 'A-1', '00021', '33', 'BORDEAUX', 'Rôle général', '15/12/2024')
    first_row = s.register_notice(first['id'], d.siren, first_notice)
    source = tmp_path / 'avis-enrichissable.pdf'
    pdf = canvas.Canvas(str(source))
    for text, y in [('MONTANT DE VOTRE IMPOT 1 234,00 EUR', 730), ('MONTANT A PAYER 934,00 EUR', 710)]:
        pdf.drawString(50, y, text)
    pdf.showPage()
    pdf.save()
    first_hash = classify_pdf(source, Path(first_row['path']), d, first_notice)
    s.notice_result(first_row['id'], 'TERMINE', '', first_hash)

    second_notice = Notice(2025, 'B-1', '00021', '33', 'BORDEAUX', 'Rôle général', '15/12/2025')
    second_row = s.register_notice(second['id'], d.siren, second_notice)
    source = tmp_path / 'avis-a-alterer.pdf'
    sample_pdf(source)
    second_hash = classify_pdf(source, Path(second_row['path']), d, second_notice)
    s.notice_result(second_row['id'], 'TERMINE', '', second_hash)
    Path(second_row['path']).write_bytes(Path(second_row['path']).read_bytes() + b'alteration')
    first_before = s.notices(first['id'])
    second_before = s.notices(second['id'])
    dossiers_before = [s.dossiers(first['id']), s.dossiers(second['id'])]
    workbook = tmp_path / 'sortie' / 'Léa' / f'Suivi Campagne {batch} — Léa.xlsx'

    with pytest.raises(ValueError, match='000000001'):
        export_batch_actions(s, batch)

    assert s.notices(first['id']) == first_before and s.notices(second['id']) == second_before
    assert [s.dossiers(first['id']), s.dossiers(second['id'])] == dossiers_before
    assert first_before[0]['extraction_status'] == ''
    assert not workbook.exists()
    assert not list(workbook.parent.glob('.collecte-*'))
