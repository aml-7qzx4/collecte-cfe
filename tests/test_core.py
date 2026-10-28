import pytest
from openpyxl import Workbook

from collecte_cfe.domain import Notice, safe_name
from collecte_cfe.portfolio import import_portfolio


def test_csv_preserves_leading_zero_and_rejects_duplicate(tmp_path):
    p = tmp_path / 'p.csv'
    p.write_text('Nom du dossier;SIREN;Collaborateur\nSAS Exemple;012345678;Léa\n', encoding='utf-8-sig')
    dossiers = import_portfolio(p)
    assert dossiers[0].siren == '012345678'
    p.write_text(p.read_text(encoding='utf-8-sig') + 'Autre;012345678;Léa\n', encoding='utf-8-sig')
    with pytest.raises(ValueError, match='doublon'):
        import_portfolio(p)


def test_xlsx_numeric_siren_and_missing_columns(tmp_path):
    w = Workbook()
    s = w.active
    s.append(['Nom du dossier', 'SIREN', 'Collaborateur'])
    s.append(['Test', 1, 'Paul'])
    p = tmp_path / 'p.xlsx'
    w.save(p)
    assert import_portfolio(p)[0].siren == '000000001'
    s.cell(1, 3, 'Autre')
    w.save(p)
    with pytest.raises(ValueError, match='Collaborateur'):
        import_portfolio(p)


def test_optional_client_email_is_imported_and_invalid_address_is_reported(tmp_path):
    p = tmp_path / 'clients.csv'
    p.write_text(
        'Nom du dossier;SIREN;Collaborateur;E-mail du client\nExemple;000000001;Léa;client@example.fr\n',
        encoding='utf-8-sig',
    )
    assert import_portfolio(p)[0].email == 'client@example.fr'
    p.write_text(
        p.read_text(encoding='utf-8-sig').replace('client@example.fr', 'client.example.fr'), encoding='utf-8-sig'
    )
    with pytest.raises(ValueError, match='e-mail'):
        import_portfolio(p)


def test_headerless_three_column_xlsx_uses_the_documented_column_order(tmp_path):
    w = Workbook()
    s = w.active
    s.append(['Alpha', 1, 'Léa'])
    s.append(['Beta', '012345678', 'Paul'])
    p = tmp_path / 'sans-entetes.xlsx'
    w.save(p)
    rows = import_portfolio(p)
    assert [(row.name, row.siren, row.collaborator) for row in rows] == [
        ('Alpha', '000000001', 'Léa'),
        ('Beta', '012345678', 'Paul'),
    ]


def test_headerless_xlsx_reports_all_rows_with_a_missing_siren(tmp_path):
    w = Workbook()
    s = w.active
    s.append(['Alpha', 1, 'Léa'])
    s.append(['Sans Siren 1', None, 'Léa'])
    s.append(['Sans Siren 2', None, 'Paul'])
    p = tmp_path / 'incomplet.xlsx'
    w.save(p)
    with pytest.raises(ValueError, match=r'Lignes 2, 3 : SIREN manquant ou invalide'):
        import_portfolio(p)


def test_the_portfolio_check_lists_every_problem_by_line_and_keeps_the_valid_dossiers(tmp_path):
    from collecte_cfe.portfolio import check_portfolio

    p = tmp_path / 'portefeuille.csv'
    p.write_text(
        'Nom du dossier;SIREN;Collaborateur\n'
        'Alpha;123456782;Léa\n'
        'Beta;41209;Léa\n'
        'Gamma;000000117;\n'
        'Alpha bis;123456782;Paul\n'
        'Delta;123456789;Paul\n'
        'A/B;000000109;Paul\n'
        'A:B;000000125;Paul\n',
        encoding='utf-8-sig',
    )

    check = check_portfolio(p)

    assert [d.name for d in check.dossiers] == ['Alpha', 'Delta', 'A/B']
    assert [(issue.line, issue.level, issue.message) for issue in check.issues] == [
        (3, 'error', 'SIREN manquant ou invalide (9 chiffres requis)'),
        (4, 'error', 'collaborateur manquant'),
        (5, 'error', 'SIREN en doublon (123456782), déjà présent ligne 2'),
        (6, 'warning', 'SIREN 123456789 à vérifier : la clé de contrôle ne correspond pas'),
        (8, 'error', 'même chemin de classement que la ligne 7'),
    ]
    assert [issue.line for issue in check.ignored] == [3, 4, 5, 8]
    assert [issue.line for issue in check.warnings] == [6]
    with pytest.raises(ValueError):
        import_portfolio(p)


def test_a_wrong_siren_check_digit_does_not_stop_a_strict_import(tmp_path):
    p = tmp_path / 'p.csv'
    p.write_text('Nom du dossier;SIREN;Collaborateur\nAlpha;000000001;Léa\n', encoding='utf-8-sig')
    assert [d.siren for d in import_portfolio(p)] == ['000000001']


def test_windows_paths_cannot_escape_or_collide(tmp_path):
    assert safe_name('../CON:') not in ('..', 'CON', '')
    assert '/' not in safe_name('../x/y')
    p = tmp_path / 'p.csv'
    p.write_text('Nom du dossier,SIREN,Collaborateur\nA/B,123456789,Paul\nA:B,987654321,Paul\n', encoding='utf-8')
    with pytest.raises(ValueError, match='classement'):
        import_portfolio(p)


@pytest.mark.parametrize(
    'role,kind',
    [
        ('Rôle d’acompte n°002', 'ACOMPTE'),
        ('Rôle général n°092', 'SOLDE'),
        ('Rôle supplémentaire n°001', 'SUPPLEMENTAIRE'),
    ],
)
def test_notice_role_kind(role, kind):
    assert Notice(2025, '123', '00021', '33', 'BORDEAUX', role, '15/12/2025').kind == kind


def test_unknown_role_is_not_a_solde():
    with pytest.raises(ValueError):
        Notice(2025, '123', '00021', '33', 'BORDEAUX', 'Document inconnu', '').kind


def test_client_emails_come_with_the_portfolio_under_their_usual_names(tmp_path):
    p = tmp_path / 'clients.csv'
    p.write_text(
        'Nom du dossier;SIREN;Collaborateur;Courriel\nAlpha;123456782;Léa;alpha@example.fr\nBeta;000000109;Léa;\n',
        encoding='utf-8-sig',
    )
    assert [(d.name, d.email) for d in import_portfolio(p)] == [('Alpha', 'alpha@example.fr'), ('Beta', '')]
    p.write_text(
        'Dossier;SIREN;Collaborateur;Adresse mail\n'.replace('Dossier', 'Nom du dossier')
        + 'Alpha;123456782;Léa;alpha@example.fr\n',
        encoding='utf-8-sig',
    )
    assert import_portfolio(p)[0].email == 'alpha@example.fr'


def test_a_headerless_portfolio_can_carry_the_email_as_a_fourth_column(tmp_path):
    w = Workbook()
    s = w.active
    s.append(['Alpha', '123456782', 'Léa', 'alpha@example.fr'])
    s.append(['Beta', '000000109', 'Paul', None])
    p = tmp_path / 'sans-entetes.xlsx'
    w.save(p)
    assert [(d.name, d.email) for d in import_portfolio(p)] == [('Alpha', 'alpha@example.fr'), ('Beta', '')]
    s.cell(1, 4, 'Note libre')
    w.save(p)
    with pytest.raises(ValueError, match='Colonnes obligatoires'):
        import_portfolio(p)
