import asyncio
import threading

from openpyxl import load_workbook

from collecte_cfe.app import MainWindow
from collecte_cfe.domain import Dossier
from collecte_cfe.situations import dossier_situation
from collecte_cfe.store import Store
from tests.test_dashboard import MemorySettings


def failed_twice(tmp_path):
    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch(
        [2025],
        tmp_path / 'out',
        [Dossier('Alpha', '000000001', 'Léa'), Dossier('Beta', '000000002', 'Léa')],
        1,
        'CFE test',
    )
    campaign = store.batch_campaigns(batch)[0]['id']
    for _attempt in range(2):
        store.dossier_result(campaign, '000000001', 'ERREUR', 'Délai dépassé')
    store.dossier_result(campaign, '000000002', 'ERREUR', 'Délai dépassé')
    return store, batch, campaign


def test_two_failures_ask_for_a_manual_check_and_a_success_starts_again(tmp_path):
    store, batch, campaign = failed_twice(tmp_path)
    alpha, beta = store.dossiers(campaign)
    assert (alpha['failures'], dossier_situation(alpha, [])) == (2, 'MANUAL')
    assert (beta['failures'], dossier_situation(beta, [])) == (1, 'RETRY'), 'one failure is still retried'
    store.verify_dossier(campaign, '000000001', True)
    assert dossier_situation(store.dossiers(campaign)[0], []) == 'CHECKED'
    store.dossier_result(campaign, '000000001', 'TERMINE')
    alpha = store.dossiers(campaign)[0]
    assert (alpha['failures'], alpha['verified_at']) == (0, '')


def test_a_dossier_to_check_by_hand_is_no_longer_relaunched_on_its_own(tmp_path):
    from collecte_cfe.sessions import run_batch
    from tests.test_sessions import IndependentPortal

    store, batch, campaign = failed_twice(tmp_path)
    portal = IndependentPortal({'now': 0, 'max': 0})
    asyncio.run(run_batch(store, batch, [portal], threading.Event(), lambda _: None, lambda *_: None, True))
    assert portal.calls == [('000000002', 2025)], 'only the dossier failed once is relaunched'
    chosen = IndependentPortal({'now': 0, 'max': 0})
    asyncio.run(
        run_batch(
            store, batch, [chosen], threading.Event(), lambda _: None, lambda *_: None, True, {(campaign, '000000001')}
        )
    )
    assert chosen.calls == [('000000001', 2025)], 'a relaunch chosen in its sheet still tries it'


def test_the_sheet_offers_to_tick_a_manual_check(tmp_path, qapp, monkeypatch):
    store, batch, campaign = failed_twice(tmp_path)
    window = MainWindow(store, settings=MemorySettings())
    window.open_campaign(batch)
    monkeypatch.setattr(window.toast, 'announce', lambda *_args, **_kwargs: None)
    assert window.result_filter_buttons['TODO'].count == 2
    rows = {window.table.item(row, 0).text(): row for row in range(window.table.rowCount())}
    assert window.table.item(rows['Alpha'], 5).text() == 'À vérifier manuellement'
    assert window.retry_button.text() == 'Relancer le dossier en erreur', 'the dossier to check by hand is not counted'
    window.table.selectRow(rows['Beta'])
    assert window.result_verified_check.isHidden()
    window.table.selectRow(rows['Alpha'])
    assert not window.result_verified_check.isHidden() and not window.result_verified_check.isChecked()
    window.result_verified_check.click()
    assert store.dossiers(campaign)[0]['verified_at']
    rows = {window.table.item(row, 0).text(): row for row in range(window.table.rowCount())}
    assert window.table.item(rows['Alpha'], 5).text() == 'Vérifié manuellement'
    assert window.result_verified_check.isChecked() and window.result_verified_check.text().startswith(
        'Vérifié manuellement le '
    )
    window.close()


def test_the_workbook_lists_the_dossiers_to_check_by_hand(tmp_path):
    from collecte_cfe.documents import export_batch_actions

    store, batch, campaign = failed_twice(tmp_path)
    workbook = load_workbook(export_batch_actions(store, batch)[0])
    sheet = workbook['Dossiers']
    headers = [cell.value for cell in sheet[6]]
    alpha = dict(zip(headers, [cell.value for cell in sheet[7]]))
    assert alpha['Dossier'] == 'Alpha' and alpha['Situation'] == 'À vérifier manuellement', 'it comes first'
    assert alpha['Information'].startswith('2 échecs : à vérifier sur impots.gouv.fr.')
    summary = [cell.value for row in workbook['Synthèse'].iter_rows() for cell in row if cell.value]
    assert '1 dossier à vérifier manuellement' in summary
    workbook.close()


def test_a_finished_campaign_is_archived_in_one_workbook(tmp_path, qapp, monkeypatch):
    store = Store(tmp_path / 'db.sqlite')
    first = store.create_batch(
        [2025],
        tmp_path / 'out',
        [Dossier('Alpha', '000000001', 'Léa'), Dossier('Beta', '000000002', 'Hugo')],
        1,
        'CFE finie',
    )
    second = store.create_batch([2025], tmp_path / 'out2', [Dossier('Gamma', '000000003', 'Léa')], 1, 'CFE en cours')
    campaign = store.batch_campaigns(first)[0]['id']
    window = MainWindow(store, settings=MemorySettings())
    announced = []
    monkeypatch.setattr(window.toast, 'announce', lambda text, target=None, **_kwargs: announced.append((text, target)))
    problems = []
    window.problem = problems.append
    assert not window.archive_campaign(first) and 'Terminez la collecte' in problems[0], (
        'only a finished campaign is archived'
    )
    for siren in ('000000001', '000000002'):
        store.dossier_result(campaign, siren, 'ERREUR', 'x')
    window.open_campaign(first)
    window.campaign_archive_action.trigger()
    path = store.batch(first)['archive_path']
    assert path.endswith('.xlsx') and (tmp_path / 'out').resolve() == __import__('pathlib').Path(path).parent.resolve()
    assert announced[-1][0] == 'Campagne archivée.'
    workbook = load_workbook(path)
    assert sorted(row[3] for row in workbook['Dossiers'].iter_rows(min_row=7, values_only=True)) == ['Hugo', 'Léa']
    workbook.close()
    assert sorted(window.campaign_nav_items) == [second], 'an archived campaign leaves the sidebar'
    assert window.pages.currentIndex() == 0
    card = next(card for card in window.campaign_cards if card.open_button.text() == 'Ouvrir l’archive')
    assert card.property('archived') and 'Sortir des archives' in [action.text() for action in card.menu.actions()]
    window.unarchive_campaign(first)
    assert sorted(window.campaign_nav_items) == sorted([first, second]) and not store.batch(first)['archived_at']
    window.close()
