from openpyxl import load_workbook

from collecte_cfe.domain import Dossier
from collecte_cfe.store import Store


def sheet_rows(sheet):
    """The table under the heading of a sheet, as dictionaries keyed by its headers."""
    headers = [cell.value for cell in sheet[6]]
    return [dict(zip(headers, [cell.value for cell in row])) for row in sheet.iter_rows(min_row=7) if row[1].value]


def test_batch_exports_one_workbook_per_collaborator_with_every_year(tmp_path):
    from collecte_cfe.documents import export_batch_actions

    s = Store(tmp_path / 'db.sqlite')
    b = s.create_batch([2024, 2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 2, 'CFE Alpha')
    paths = export_batch_actions(s, b)
    assert [path.name for path in paths] == ['Suivi CFE Alpha — Léa.xlsx'] and paths[0].parent.name == 'Léa'
    workbook = load_workbook(paths[0])
    assert workbook.sheetnames == ['Synthèse', 'Dossiers', 'Avis']
    rows = sheet_rows(workbook['Dossiers'])
    assert [(row['Dossier'], row['Année'], row['Situation']) for row in rows] == [
        ('Alpha', 2024, 'En attente'),
        ('Alpha', 2025, 'En attente'),
    ]
    assert workbook['Dossiers'].freeze_panes == 'C7' and workbook['Dossiers'].tables['TableauDossiers'].ref == 'A6:P8'
    assert workbook['Dossiers'].auto_filter.ref is None, (
        'only the table filters: a second filter makes Excel refuse the file'
    )
    summary = [cell.value for row in workbook['Synthèse'].iter_rows() for cell in row if cell.value is not None]
    assert 'Par année' in summary and 'Par collaborateur' not in summary
    workbook.close()


def test_a_cabinet_workbook_gathers_every_collaborator(tmp_path):
    from collecte_cfe.documents import export_batch_actions

    s = Store(tmp_path / 'db.sqlite')
    b = s.create_batch(
        [2025],
        tmp_path / 'out',
        [Dossier('Alpha', '000000001', 'Léa'), Dossier('Beta', '000000002', 'Hugo')],
        1,
        'CFE test',
    )
    paths = export_batch_actions(s, b)
    assert [path.relative_to(tmp_path / 'out').as_posix() for path in paths] == [
        'Hugo/Suivi CFE test — Hugo.xlsx',
        'Léa/Suivi CFE test — Léa.xlsx',
        'Suivi CFE test — cabinet.xlsx',
    ]
    workbook = load_workbook(paths[-1])
    assert sorted(row['Collaborateur'] for row in sheet_rows(workbook['Dossiers'])) == ['Hugo', 'Léa']
    summary = workbook['Synthèse']
    assert summary['B4'].value == 'Cabinet · 2 collaborateurs'
    texts = [cell.value for row in summary.iter_rows() for cell in row if cell.value is not None]
    assert 'Par collaborateur' in texts and 'Hugo' in texts and 'Léa' in texts
    workbook.close()


def test_the_summary_says_where_the_dossiers_stand_and_what_to_do(tmp_path):
    from collecte_cfe.documents import export_batch_actions, proof_pdf

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch(
        [2025],
        tmp_path / 'out',
        [
            Dossier('Erreur', '000000001', 'Léa'),
            Dossier('Sans accès', '000000004', 'Léa'),
            Dossier('Sans avis', '000000005', 'Léa'),
        ],
        1,
    )
    campaign = store.batch_campaigns(batch)[0]['id']
    store.dossier_result(campaign, '000000001', 'ERREUR', 'Téléchargement interrompu')
    store.dossier_result(campaign, '000000004', 'SANS_HABILITATION')
    proof = store.register_proof(campaign, '000000005')
    store.proof_result(campaign, '000000005', proof_pdf(proof, Dossier('Sans avis', '000000005', 'Léa'), 2025))
    store.dossier_result(campaign, '000000005', 'ABSENT')

    workbook = load_workbook(export_batch_actions(store, batch)[0])
    summary = workbook['Synthèse']
    lines = {summary.cell(row, 2).value: summary.cell(row, 3).value for row in range(8, 12)}
    assert lines == {'À relancer ou sans habilitation': 2, 'Sans avis': 1, 'Total': 3, None: None}
    assert summary['B8'].hyperlink.location == "'Dossiers'!A6"
    todo = [summary.cell(row, 2).value for row in range(15, 25) if summary.cell(row, 2).value]
    assert todo[:2] == ['À faire', '2 dossiers à relancer ou sans habilitation']
    rows = sheet_rows(workbook['Dossiers'])
    assert [row['Situation'] for row in rows] == ['À relancer', 'Sans habilitation', 'Aucun avis'], (
        'what is left to do comes first'
    )
    assert rows[0]['Information'] == 'Téléchargement interrompu'
    assert workbook['Dossiers'].cell(9, 15).hyperlink is not None, 'the absence proof opens from its row'
    workbook.close()


def test_dossier_status_counts_reconcile_with_total():
    from collecte_cfe.domain import summarize_dossiers

    rows = [
        {'status': status} for status in ['TERMINE', 'ABSENT', 'SANS_HABILITATION', 'ERREUR', 'EN_ATTENTE', 'EN_COURS']
    ]
    metrics = summarize_dossiers(rows)
    assert metrics['total'] == 6
    assert sum(metrics[key] for key in ('finished', 'absent', 'no_access', 'retry', 'pending')) == 6
    assert metrics['absent'] == 1 and metrics['pending'] == 2


def test_two_notices_for_one_dossier_make_two_detail_rows_but_one_dossier(tmp_path):
    from pathlib import Path

    from reportlab.pdfgen import canvas

    from collecte_cfe.documents import classify_pdf, export_batch_actions
    from collecte_cfe.domain import Notice
    from collecte_cfe.notice_extraction import NoticeExtraction

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('=Alpha', '000000001', 'Léa')], 1)
    campaign = store.batch_campaigns(batch)[0]['id']
    for role, reference, amount in [('Rôle d’acompte', 'A-1', 50000), ('Rôle général', 'S-1', 95000)]:
        notice = store.register_notice(
            campaign, '000000001', Notice(2025, reference, '00021', '33', 'Bordeaux', role, '15/12/2025')
        )
        source = tmp_path / f'{reference}.pdf'
        pdf = canvas.Canvas(str(source))
        pdf.drawString(50, 700, 'Avis synthétique')
        pdf.showPage()
        pdf.save()
        digest = classify_pdf(
            source,
            Path(notice['path']),
            Dossier('=Alpha', '000000001', 'Léa'),
            Notice(2025, reference, '00021', '33', 'Bordeaux', role, '15/12/2025'),
        )
        store.notice_result(notice['id'], 'TERMINE', '', digest)
        store.notice_extraction_result(
            notice['id'],
            NoticeExtraction(
                tax_amount_cents=amount,
                amount_due_cents=amount,
                payment_method='PRELEVEMENT_ECHEANCE',
                bank_account_display='FR** 1234',
                bank_account_masked=True,
                review_status='PARTIEL',
            ),
        )
    store.dossier_result(campaign, '000000001', 'TERMINE')
    workbook = load_workbook(export_batch_actions(store, batch)[0])
    sheet = workbook['Avis']
    headers = [cell.value for cell in sheet[6]]
    assert sheet.max_row == 8
    assert headers == [
        'Dossier',
        'SIREN',
        'Année',
        'Type',
        'Référence',
        'Montant de l’avis',
        'À payer',
        'Mode de paiement',
        'Échéance',
        'Commune',
        'Lecture',
        'PDF',
    ]
    assert {sheet.cell(row, headers.index('Type') + 1).value for row in (7, 8)} == {'Acompte', 'Solde'}
    assert sheet['A7'].data_type == 's' and sheet['A7'].value == '=Alpha'
    assert sheet.freeze_panes == 'C7' and sheet.tables['TableauAvis'].ref == 'A6:L8'
    assert sheet.cell(7, headers.index('Échéance') + 1).is_date
    rows = sheet_rows(workbook['Dossiers'])
    assert len(rows) == 1 and rows[0]['À payer'] == 1450 and rows[0]['Situation'] == 'Prélevé'
    assert rows[0]['Compte prélevé'] == 'FR** 1234', 'one account for both debited notices'
    assert rows[0]['Avis'] == '2 avis', 'two notices open their folder'
    workbook.close()


def test_failed_detected_notice_is_visible_without_a_classified_pdf(tmp_path):
    from collecte_cfe.documents import export_batch_actions
    from collecte_cfe.domain import Notice
    from collecte_cfe.notice_extraction import NoticeExtraction

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 1)
    campaign = store.batch_campaigns(batch)[0]['id']
    notice = store.register_notice(
        campaign, '000000001', Notice(2025, 'A-1', '00021', '33', 'Bordeaux', 'Rôle d’acompte', '15/06/2025')
    )
    store.notice_extraction_result(
        notice['id'],
        NoticeExtraction(
            tax_amount_cents=50000,
            amount_due_cents=50000,
            payment_method='PRELEVEMENT_ECHEANCE',
            bank_account_display='FR** 1234',
            bank_account_masked=True,
            review_status='PARTIEL',
        ),
    )
    store.notice_result(notice['id'], 'ERREUR', 'Délai dépassé lors du téléchargement')
    store.dossier_result(campaign, '000000001', 'ERREUR', 'Délai dépassé lors du téléchargement')
    workbook = load_workbook(export_batch_actions(store, batch)[0])
    (values,) = sheet_rows(workbook['Avis'])
    assert values['Lecture'] == 'À relancer · Délai dépassé lors du téléchargement'
    assert not values['PDF'] and values['À payer'] is None and values['Mode de paiement'] in ('', None)
    (row,) = sheet_rows(workbook['Dossiers'])
    assert row['Situation'] == 'À relancer' and row['À payer'] is None and not row['Compte prélevé']
    workbook.close()


def test_batch_exports_a_french_csv_with_notice_durations(tmp_path):
    import csv

    from collecte_cfe.documents import export_batch_timings
    from collecte_cfe.domain import Notice

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 1, name='Test réel')
    campaign = store.batch_campaigns(batch)[0]['id']
    store.record_timing(
        batch,
        campaign,
        '000000001',
        1,
        'AVIS',
        'TERMINE',
        2.543,
        notice=Notice(2025, 'REF-42', '00021', '33', 'Bordeaux', 'Rôle d’acompte', ''),
    )
    target = export_batch_timings(store, batch)
    assert target.name == 'Temps_collecte_CFE_Test réel.csv'
    with target.open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.DictReader(stream, delimiter=';'))
    assert rows[0]['Session'] == '1' and rows[0]['Type d’avis'] == 'ACOMPTE'
    assert rows[0]['Durée (secondes)'] == '2,543'


def test_a_refund_read_on_a_notice_is_stored(tmp_path):
    from collecte_cfe.domain import Notice
    from collecte_cfe.notice_extraction import NoticeExtraction

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 1)
    campaign = store.batch_campaigns(batch)[0]['id']
    notice = store.register_notice(
        campaign, '000000001', Notice(2025, 'S-1', '00021', '33', 'Bordeaux', 'Rôle général', '15/12/2025')
    )
    store.notice_extraction_result(
        notice['id'], NoticeExtraction(30700, 0, 'AUCUN_PAIEMENT', None, False, 'EXTRAIT', 2100)
    )
    assert store.notices(campaign)[0]['refund_cents'] == 2100


def test_notices_classified_by_hand_are_not_counted_as_readings_to_check(tmp_path):
    from pathlib import Path

    from reportlab.pdfgen import canvas

    from collecte_cfe.documents import classify_pdf, export_batch_actions
    from collecte_cfe.domain import Notice
    from collecte_cfe.notice_extraction import NoticeExtraction

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch(
        [2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa'), Dossier('Beta', '000000002', 'Léa')], 1
    )
    campaign = store.batch_campaigns(batch)[0]['id']
    for siren, name, status in [('000000001', 'Alpha', 'MANUEL'), ('000000002', 'Beta', 'A_VERIFIER')]:
        data = Notice(2025, 'S-' + siren, '00021', '33', 'Bordeaux', 'Rôle général', '15/12/2025')
        notice = store.register_notice(campaign, siren, data)
        source = tmp_path / f'{siren}.pdf'
        pdf = canvas.Canvas(str(source))
        pdf.drawString(50, 700, 'Avis synthétique')
        pdf.showPage()
        pdf.save()
        digest = classify_pdf(source, Path(notice['path']), Dossier(name, siren, 'Léa'), data)
        store.notice_result(notice['id'], 'TERMINE', '', digest)
        store.notice_extraction_result(notice['id'], NoticeExtraction(50000, 40000, 'A_REGLER', None, False, status))
        store.dossier_result(campaign, siren, 'TERMINE')
    workbook = load_workbook(export_batch_actions(store, batch)[0])
    assert sorted(row['Lecture'] for row in sheet_rows(workbook['Avis'])) == ['Classée à la main', 'À vérifier']
    assert {row['Dossier']: row['Situation'] for row in sheet_rows(workbook['Dossiers'])} == {
        'Alpha': 'À régler',
        'Beta': 'À vérifier',
    }
    summary = workbook['Synthèse']
    assert (summary['B8'].value, summary['C8'].value) == ('Lecture à vérifier', 1)
    assert (summary['B9'].value, summary['C9'].value) == ('À régler par le client', 1)


def test_only_dossiers_with_classified_notices_can_be_pointed(tmp_path):
    import pytest

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch(
        [2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa'), Dossier('Beta', '000000002', 'Léa')], 1
    )
    campaign = store.batch_campaigns(batch)[0]['id']
    store.dossier_result(campaign, '000000001', 'TERMINE')

    stamp = store.point_dossier(campaign, '000000001', True)

    alpha, beta = store.dossiers(campaign)
    assert alpha['pointed_at'] == stamp and len(stamp) == len('2025-12-01T10:41:00')
    assert beta['pointed_at'] == ''
    with pytest.raises(ValueError, match='avis classés'):
        store.point_dossier(campaign, '000000002', True)
    assert store.point_dossier(campaign, '000000001', False) == ''
    assert store.dossiers(campaign)[0]['pointed_at'] == ''


def test_the_follow_up_sheet_shows_when_a_dossier_was_pointed(tmp_path):
    from datetime import datetime

    from collecte_cfe.documents import export_batch_actions

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 1)
    campaign = store.batch_campaigns(batch)[0]['id']
    store.dossier_result(campaign, '000000001', 'TERMINE')
    store.point_dossier(campaign, '000000001', True)

    sheet = load_workbook(export_batch_actions(store, batch)[0])['Dossiers']

    headers = [cell.value for cell in sheet[6]]
    assert headers[0] == 'Pointé le'
    assert isinstance(sheet.cell(7, 1).value, datetime)


def test_the_dossiers_sheet_explains_the_change_since_last_year(tmp_path):
    from pathlib import Path

    from collecte_cfe.cfe_details import read_details
    from collecte_cfe.documents import classify_pdf, export_batch_actions
    from collecte_cfe.domain import Notice
    from collecte_cfe.notice_extraction import NoticeExtraction
    from tests.test_cfe_details import write_notice

    store = Store(tmp_path / 'db.sqlite')
    dossier = Dossier('SAS Exemple', '000000001', 'Léa')
    batch = store.create_batch([2025], tmp_path / 'out', [dossier], 1)
    campaign = store.batch_campaigns(batch)[0]['id']
    data = Notice(2025, 'REF-2025', '00018', '75', 'PARIS', 'Rôle général n°092', '15/12/2025')
    row = store.register_notice(campaign, dossier.siren, data)
    source, total = write_notice(tmp_path, 'n.pdf', year=2025)
    store.notice_result(row['id'], 'TERMINE', '', classify_pdf(source, Path(row['path']), dossier, data))
    store.notice_extraction_result(
        row['id'], NoticeExtraction(total * 100, total * 100, 'A_REGLER', None, False, 'EXTRAIT')
    )
    store.dossier_result(campaign, dossier.siren, 'TERMINE')
    previous, _ = write_notice(tmp_path, 'n-1.pdf', year=2024, rate='26,43', previous_rate=None, fees=100)
    store.save_breakdown(
        dossier.siren,
        Notice(2024, 'REF-2024', '00018', '75', 'PARIS', 'Rôle général n°092', ''),
        read_details(previous),
        'ANNEE_PRECEDENTE',
    )
    workbook = load_workbook(export_batch_actions(store, batch)[0])
    (row,) = sheet_rows(workbook['Dossiers'])
    assert row['Évolution sur un an'].startswith('Hausse de ') and '• ' in row['Évolution sur un an']
    assert row['Situation'] == 'À régler'
    workbook.close()
