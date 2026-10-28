import asyncio
import threading
from pathlib import Path

from PIL import Image
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from collecte_cfe.domain import Dossier, Notice
from collecte_cfe.notice_extraction import NoticeExtraction, debited_account, extract_notice_pdf
from collecte_cfe.store import Store
from tests.test_sessions import ScopedPortal


def test_only_direct_debits_give_the_debited_account():
    debit = {
        'status': 'TERMINE',
        'payment_method': 'PRELEVEMENT_ECHEANCE',
        'bank_account_display': 'FR76 3000 XXXX 0123',
    }
    monthly = {
        'status': 'TERMINE',
        'payment_method': 'PRELEVEMENT_MENSUEL',
        'bank_account_display': 'FR76 1111 XXXX 0456',
    }
    printed_for_refund = {
        'status': 'TERMINE',
        'payment_method': 'AUCUN_PAIEMENT',
        'bank_account_display': 'FR76 9999 XXXX 0789',
    }
    not_filed = {**debit, 'status': 'ERREUR', 'bank_account_display': 'FR76 2222 XXXX 0000'}
    assert (
        debited_account([debit, debit, monthly, printed_for_refund, not_filed])
        == 'FR76 3000 XXXX 0123 · FR76 1111 XXXX 0456'
    )
    assert debited_account([]) == ''


def test_the_campaign_table_shows_the_debited_account(tmp_path, qapp):
    from collecte_cfe.app import MainWindow
    from collecte_cfe.ui_results import COL_ACCOUNT
    from tests.test_dashboard import MemorySettings

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 1, 'CFE')
    campaign = store.batch_campaigns(batch)[0]['id']
    row = store.register_notice(
        campaign, '000000001', Notice(2025, 'R1', '00018', '75', 'PARIS', 'Rôle général', '15/12/2025')
    )
    store.notice_result(row['id'], 'TERMINE', '', 'sha')
    store.notice_extraction_result(
        row['id'], NoticeExtraction(50000, 50000, 'PRELEVEMENT_ECHEANCE', 'FR76 3000 XXXX 0123', True, 'PARTIEL')
    )
    store.dossier_result(campaign, '000000001', 'TERMINE')
    window = MainWindow(store, settings=MemorySettings())
    window.open_campaign(batch)
    assert window.table.horizontalHeaderItem(COL_ACCOUNT).text() == 'Compte prélevé'
    assert window.table.item(0, COL_ACCOUNT).text() == 'FR76 3000 XXXX 0123'
    window.close()


def notice_page(path, *, cover):
    pdf = canvas.Canvas(str(path))
    if cover:
        pdf.drawString(50, 700, 'Identification de l’avis')
        pdf.showPage()
    for y, line in (
        (730, 'MONTANT DE VOTRE IMPOT 1 234,00 EUR'),
        (710, 'MONTANT A PAYER 934,00 EUR'),
        (690, "Vous avez choisi le PRELEVEMENT A L'ECHEANCE."),
        (670, 'sur le compte bancaire FR00 0000 0000 0000 XXXXXXXX 000.'),
    ):
        pdf.drawString(50, y, line)
    pdf.showPage()
    pdf.save()


def test_notices_are_read_with_or_without_the_old_cover_page(tmp_path):
    from pypdf import PdfWriter

    from collecte_cfe.documents import classify_pdf

    d = Dossier('Exemple', '000000001', 'Léa')
    n = Notice(2025, '123', '00021', '33', 'BORDEAUX', 'Rôle général', '')
    source = tmp_path / 'portail.pdf'
    notice_page(source, cover=False)
    filed = tmp_path / 'nouveau.pdf'
    classify_pdf(source, filed, d, n)
    assert extract_notice_pdf(filed).tax_amount_cents == 123400
    old = tmp_path / 'ancien.pdf'
    notice_page(old, cover=True)
    writer = PdfWriter()
    writer.append(PdfReader(old))
    writer.add_metadata({'/CollecteNoticeKey': n.key})
    writer.write(old)
    assert extract_notice_pdf(old).tax_amount_cents == 123400, (
        'notices filed by the previous version keep their cover page'
    )


class CapturingPortal(ScopedPortal):
    async def capture_search(self, siren, path):
        Image.new('RGB', (400, 200), 'white').save(path)
        self.captured = siren
        return True


def test_a_dossier_without_access_keeps_each_accounts_search_in_its_folder(tmp_path):
    from openpyxl import load_workbook

    from collecte_cfe.documents import export_batch_actions
    from collecte_cfe.sessions import run_batch

    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 2)
    account = store.add_expert_account('Cabinet', 'expert@example.fr')
    store.assign_expert_account(0, account)
    active = {'now': 0, 'max': 0}
    portals = [CapturingPortal(active, {'000000001': 'Alpha'}, denied={'000000001'}) for _ in range(2)]
    asyncio.run(run_batch(store, batch, portals, threading.Event(), lambda _: None, lambda *_: None))
    campaign = store.batch_campaigns(batch)[0]['id']
    row = store.dossiers(campaign)[0]
    assert row['status'] == 'SANS_HABILITATION'
    proof = Path(row['proof_path'])
    assert proof.name == 'Alpha_CFE_2025_SANS_HABILITATION.pdf' and proof.parent == tmp_path / 'out' / 'Léa' / 'Alpha'
    reader = PdfReader(proof)
    assert len(reader.pages) == 2 and reader.metadata['/CollecteProof'] == 'SANS_HABILITATION'
    assert (
        'compte Cabinet (expert@example.fr)' in reader.pages[0].extract_text()
        and 'session 2' in reader.pages[1].extract_text()
    )
    assert row['proof_sha256']
    sheet = load_workbook(export_batch_actions(store, batch)[0])['Dossiers']
    headers = [cell.value for cell in sheet[6]]
    assert sheet.cell(7, headers.index('Capture du portail') + 1).hyperlink is not None


def test_a_proof_of_the_other_kind_gets_its_own_file(tmp_path):
    store = Store(tmp_path / 'db.sqlite')
    batch = store.create_batch([2025], tmp_path / 'out', [Dossier('Alpha', '000000001', 'Léa')], 1)
    campaign = store.batch_campaigns(batch)[0]['id']
    no_access = store.register_proof(campaign, '000000001', 'SANS_HABILITATION')
    store.proof_result(campaign, '000000001', 'sha')
    assert store.register_proof(campaign, '000000001', 'SANS_HABILITATION') == no_access
    absence = store.register_proof(campaign, '000000001')
    assert absence.name == 'Alpha_CFE_2025_ABSENCE.pdf' and absence != no_access
    assert store.dossiers(campaign)[0]['proof_sha256'] == '', 'the new proof is not taken for the old one'


def test_a_notice_without_a_readable_year_is_confirmed_by_its_reference(tmp_path):
    """The reference printed on a CFE notice starts with its year: « 25 99 000000 101 » is a 2025 notice."""
    import pytest

    from collecte_cfe.documents import verify_original_notice

    pdf = tmp_path / 'avis.pdf'
    page = canvas.Canvas(str(pdf))
    page.drawString(
        40, 700, "Numéro fiscal : 000000133 00034   Référence de l'avis : 25 99 000000 101   Lieu d'imposition : 0001"
    )
    page.save()
    d = Dossier('EXEMPLE', '000000133', 'Léa')
    verify_original_notice(pdf, d, Notice(2025, '2599000000101', '00034', '99', 'COMMUNE', 'Rôle général n°092', ''))
    with pytest.raises(ValueError):
        verify_original_notice(
            pdf, d, Notice(2024, '2599000000101', '00034', '99', 'COMMUNE', 'Rôle général n°092', '')
        )
    with pytest.raises(ValueError):
        verify_original_notice(
            pdf, d, Notice(2025, '2599000000999', '00034', '99', 'COMMUNE', 'Rôle général n°092', '')
        )
